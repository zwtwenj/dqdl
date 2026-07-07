"""
生成基础场景设定文档
"""
import os, time
from openai import OpenAI
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(__file__), '.env'))
client = OpenAI(
    api_key=os.getenv('apikey'),
    base_url='https://api.deepseek.com'
)
RAG_DIR = os.path.join(os.path.dirname(__file__), 'rag')

def call_ai(system, user, temp=0.8, max_tokens=8000):
    r = client.chat.completions.create(
        model='deepseek-chat',
        messages=[{'role':'system','content':system},{'role':'user','content':user}],
        temperature=temp, max_tokens=max_tokens,
    )
    return r.choices[0].message.content

def save_docx(filename, title, content):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(11)
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in content.split('\n'):
        line = line.strip()
        if not line: continue
        if line.startswith('# ') or line.startswith('一') or line.startswith('二') or line.startswith('三') or line.startswith('四') or line.startswith('五') or line.startswith('六') or line.startswith('七') or line.startswith('八') or line.startswith('九') or line.startswith('十'):
            doc.add_heading(line.lstrip('# '), level=2)
        else:
            doc.add_paragraph(line)
    path = os.path.join(RAG_DIR, filename)
    doc.save(path)
    print(f'  saved: {filename} ({os.path.getsize(path):,} bytes)')

SYSTEM = """你是斗气大陆（斗破苍穹）世界观的游戏设计师。
严格遵循斗气大陆的修炼体系（斗之气→斗者→斗师→大斗师→斗灵→斗王→斗皇→斗宗→斗尊→半圣→斗圣→斗帝）。
所有设定必须符合这个世界观，货币为金币/银币/铜币，1金币=100银币=10000铜币。
输出要求结构清晰，用编号和分点格式。"""

# 场景列表
scenes = [
    {
        "name": "拍卖行",
        "prompt": """请设计「拍卖行」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：高端物品交易平台，每样拍品独立拍卖，一般有较高溢价（市场价的1.5-3倍）
- 拍卖频次：每日定时拍卖（如每天午时一场），特殊节日额外场
- 建筑外观与内部布局描述

二、拍卖品分级
- 普通场：三阶以下丹药、二阶以下功法、普通魔核
- 精品场：三阶丹药、玄阶功法、三阶魔核、异火碎片
- 压轴场：四阶以上丹药、地阶功法碎片、罕见异兽材料
- 每场的起拍价和加价幅度规则

三、拍卖机制
- 竞拍规则（起拍价、加价幅度、成交）
- 暗拍机制（部分稀有物品盲拍）
- 手续费（成交价的10%-15%）
- 保证金制度

四、捡漏规则
- 极小概率出现被错误估价的珍品（概率1%-5%）
- 触发条件（玩家智力/运气属性检定）

五、NPC设定
- 拍卖师（固定NPC，性格描写）
- 鉴定师（可付费鉴定物品真伪）
- 常见竞拍对手类型（土豪型、内行型、盲目型）"""
    },
    {
        "name": "佣兵公会",
        "prompt": """请设计「佣兵公会」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：任务发布、接取、交付的中心枢纽
- 建筑外观与内部布局描述

二、任务系统
- 任务分类：战斗型、收集型、护送型、探索型、紧急型
- 每种类型的详细说明和示例任务（3-5个具体任务）
- 任务等级：D/C/B/A/S，对应境界要求和报酬范围

三、任务流程
- 接取流程（查看→接取→执行→交付）
- 任务时限规则
- 失败惩罚（信誉度扣除、赔偿）
- 组队任务规则

四、佣兵等级与权益
- 铁级/铜级/银级/金级/暗金级
- 每级可接的任务等级上限
- 每级的专属福利（折扣、优先权、紧急救援）
- 升级条件

五、NPC设定
- 公会接待员
- 任务发布板管理员
- 佣兵训练教官（可付费学习基础战斗技巧）"""
    },
    {
        "name": "坊市",
        "prompt": """请设计「坊市」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：中低端物品交易集散地，数量多品种杂
- 环境氛围描写（嘈杂、鱼龙混杂、地摊+商铺混合）
- 建筑布局

二、交易系统
- 摊位区：随机生成商品（低价、品质一般）
- 商铺区：固定商品（价格稍高但品质有保障）
- 黑市角落：需要特定条件才能进入，出售违禁品
- 讨价还价机制（智力检定，可砍价5%-30%）

三、捡漏系统（核心）
- 极小概率（2%-8%）出现被错误定价的高端物品
- 捡漏触发条件：运气属性检定
- 捡漏物品示例（10个具体例子）
- 捡漏后的"被追回"风险（部分高价值物品可能被原主追回）

四、随机事件
- 坊市中可能遇到的随机事件（5-8种）
- 如：小偷、假货骗子、黑心商人、隐藏高手摆摊

五、NPC设定
- 常见摊贩类型（5种）
- 巡逻卫兵（治安维护）"""
    },
    {
        "name": "炼药师公会",
        "prompt": """请设计「炼药师公会」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：丹药交易、炼药学习、丹炉租赁
- 建筑外观与内部布局（丹炉房、药材库、交易厅、教学区）

二、丹药交易
- 公会出售的丹药清单（按品阶分类，每种标注价格和效果）
- 玩家出售丹药的回收价格（市价的50%-70%）
- 稀有丹药定制委托流程

三、丹炉租赁
- 丹炉等级：铁炉/铜炉/银炉/金炉
- 每级的租金（按次/按天计费）
- 每级的炼药成功率加成
- 租赁押金规则
- 炸炉赔偿条款

四、炼药学习
- 基础炼药术教学（学费、课程内容）
- 丹方购买（不同品阶丹方的价格）
- 炼药师等级考核（一品到三品）

五、NPC设定
- 公会会长（炼药师大师）
- 炼药指导师
- 药材管理员"""
    },
    {
        "name": "武技修炼场",
        "prompt": """请设计「武技修炼场」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：付费使用修炼设施，加速修炼/突破
- 建筑布局

二、修炼设施
- 基础修炼室：斗气恢复速度+20%，价格低
- 元气修炼室：斗气恢复速度+50%，适合日常修炼
- 突破密室：突破成功率+10%，价格高
- 实战模拟区：与傀儡对战，提升战斗经验
- 每种的收费标准和使用规则

三、功法教学
- 出售黄阶/玄阶功法残卷
- 功法练习场（有教官指导）
- 功法价格表

四、NPC设定
- 修炼场管理员
- 功法教官"""
    },
    {
        "name": "旅馆与休息区",
        "prompt": """请设计「旅馆与休息区」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：恢复HP/MP、存档、触发随机事件
- 不同等级旅馆的描述

二、房间类型
- 大通铺：便宜，恢复慢，可能被偷
- 普通房：适中，恢复正常
- 上房：较贵，恢复快，可能触发NPC事件
- 雅间：最贵，全恢复，高概率触发特殊事件
- 每种的恢复效率和价格

三、随机事件
- 住店时可能触发的事件（8-10种）
- 如：隔壁房间传来密谋、半夜有人敲门、梦中获得线索

四、存档机制
- 住店=存档点
- 不同房间存档额外效果"""
    },
    {
        "name": "药材商行",
        "prompt": """请设计「药材商行」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：药材收购与出售
- 建筑布局

二、药材交易
- 常见药材出售清单（按品阶分类，标注价格）
- 玩家采集药材的回收价格
- 批量交易折扣
- 稀有药材预订服务

三、代炼药服务
- 玩家提供药材+加工费→获得丹药
- 不同品阶丹药的加工费
- 成功率和失败赔偿规则

四、NPC设定
- 掌柜（懂药理，可能提供线索）
- 采药向导（可雇佣，带路去药产地）"""
    },
    {
        "name": "铁匠铺",
        "prompt": """请设计「铁匠铺」场景的完整游戏设定。要求：

一、场景概述
- 功能定位：武器防具购买、修理、定制、强化
- 建筑布局

二、装备系统
- 武器分类（刀、剑、枪、斧、尺等）和价格表
- 防具分类（皮甲、铁甲、灵甲等）和价格表
- 装备品级：普通/精良/稀有/传说
- 装备修理费用

三、强化系统
- 魔核镶嵌：将魔核镶嵌到装备上，赋予属性
- 不同等级魔核的镶嵌效果和费用
- 镶嵌成功率
- 拆除已镶嵌魔核的规则

四、定制武器
- 提供材料+设计+加工费
- 定制周期
- 特殊效果附加

五、NPC设定
- 铁匠大师
- 学徒"""
    },
]

def main():
    all_content = []
    for i, scene in enumerate(scenes):
        print(f'[{i+1}/{len(scenes)}] 生成: {scene["name"]}...')
        content = call_ai(SYSTEM, scene["prompt"], temp=0.75, max_tokens=4000)
        all_content.append(f"\n\n{'='*50}\n\n# {scene['name']}\n\n{content}")
        print(f'  done ({len(content)} chars)')
        time.sleep(1)

    full = "# 斗气大陆·基础场景设定\n\n本文档定义了游戏中所有基础场景的功能、交易系统、NPC设定和交互规则。\n" + "".join(all_content)
    save_docx("基础场景设定.docx", "基础场景设定", full)
    print('\nDone!')

if __name__ == '__main__':
    main()
