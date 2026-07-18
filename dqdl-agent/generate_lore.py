"""
一次性脚本：调用 DeepSeek 生成「剧情素材」文档，写入 rag/剧情素材.docx。

背景：现有向量库只有图鉴(魔兽/材料/丹药)和系统设定，缺乏「叙事素材」——
势力档案、NPC原型、桥段模板、地域风物。这类素材是多分支剧本生成的叙事基石，
RAG 只能基于已有内容检索，故必须先补齐。

生成 4 类素材（每类一次 LLM 调用），均要求「【名称】...【...】...」条目格式，
便于 rag_service.split_by_entry 按条目切分成 chunk。
生成后手动 review，再运行 rag_service.py index 入库。

用法：
  cd dqdl-agent
  python generate_lore.py            # 生成并写入 docx
  python generate_lore.py --dry      # 只打印不写文件
"""
import os
import sys
import json
from llm_client import call_deepseek
from utils import _parse_json_object
from docx import Document

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(BASE_DIR)
RAG_DIR = os.path.join(PROJECT_DIR, 'rag')
OUT_PATH = os.path.join(RAG_DIR, '剧情素材.docx')

# 世界观摘要（约束 LLM 不偏离斗气大陆设定；与 世界观与核心设定.docx 对齐）
WORLD_SUMMARY = """斗气大陆世界观（《斗破苍穹》）：
- 修炼体系：斗之气→斗者→斗师→大斗师→斗灵→斗王→斗皇→斗宗→斗尊→半圣→斗圣→斗帝，每阶分1-9星。
- 地理：西北地域(加玛帝国/出云帝国等，每五年宗派大会)、黑角域(三不管地带，迦南学院)、中州(核心繁华)、隐秘空间界(远古八族)。
- 势力：远古八族(古/魂/炎等，族长皆顶尖斗圣)、魂殿(反派，收魂)、丹塔(炼药师圣地)、两宗三谷四方阁(中州明面势力)。
- 职业：炼药师(火木双属性，一至九品+帝品，高阶丹药引天地异象)、佣兵、散修。
- 异火榜：23种天地异火，融合可施组合技(佛怒火莲)，反噬极高。
- 功法斗技分天地玄黄四阶；魔兽有品阶(一阶~九阶+)与属性(金木水火土风雷)。
- 氛围：强者为尊的丛林法则，丹药/魔核/异火是核心战略资源。"""


def gen_section(system_prompt, user_prompt, call_type='event'):
    """调 LLM 生成一类素材，返回纯文本（条目格式）。失败抛错。"""
    content, _ = call_deepseek(
        system_prompt, user_prompt,
        temperature=0.9, max_tokens=4000, call_type=call_type,
    )
    # 尝试剥 markdown 代码块
    s = content.strip()
    if s.startswith('```'):
        s = s.split('```')
        # 取最长的非空段（通常是正文）
        s = max((p.strip() for p in s if p.strip() and not p.lower().startswith('json')), key=len)
    return s


# ============================================================
#  四类素材的 prompt
# ============================================================

def prompt_factions():
    """势力档案：用于剧本里势力相关恩怨/委托/遭遇。"""
    sys = (
        '你是斗气大陆世界观设定师。基于给定世界观，编写「势力档案」素材。'
        '每个势力是一个独立条目，用【名称】开头，字段齐全。'
        '只输出条目文本，不要前言后语、不要 markdown 代码块。'
    )
    user = (
        f'【世界观约束】\n{WORLD_SUMMARY}\n\n'
        '【任务】编写 10 个势力的档案，覆盖：'
        '加玛帝国本土势力(如某大家族/佣兵团)、中州某宗门、黑角域某势力、魂殿外围势力、'
        '炼药师组织、散修联盟、魔兽族群附庸势力等。要有大有小(玩家前期可接触的低阶势力也要有)。\n\n'
        '【条目格式】严格如下（每条空行分隔）：\n'
        '【名称】势力名\n'
        '【类型】家族/宗门/佣兵团/帝国官方/地下势力/散修组织/魔兽势力\n'
        '【势力等级】低(斗者~斗师) / 中(大斗师~斗灵) / 高(斗王~斗宗) / 顶尖(斗尊+)\n'
        '【据点】所在地（须符合世界观地理）\n'
        '【行事风格】一句话（如"唯利是图""护短""阴险"）\n'
        '【宿敌】该势力的仇家或竞争对手（可指向另一个势力名，制造剧本冲突）\n'
        '【与玩家可能的交集】剧本切入点（如"悬赏猎杀某魔兽""争夺某药材产地""招揽玩家"）\n'
        '【代表人物原型】该势力里可能登场的一类NPC（如"铁血长老""纨绔少主"）'
    )
    return sys, user


def prompt_npcs():
    """NPC原型：剧本登场角色的性格/动机/说话方式模板。"""
    sys = (
        '你是斗气大陆世界观设定师。编写「NPC原型」素材，供剧本生成时随机抽取登场角色。'
        '每个原型是独立条目，用【名称】开头。只输出条目文本，不要前言后语、不要 markdown。'
    )
    user = (
        f'【世界观约束】\n{WORLD_SUMMARY}\n\n'
        '【任务】编写 15 个 NPC 原型，覆盖修真界各类角色：'
        '云游散修、佣兵团长、炼药师、丹徒、魔修、落魄贵族、赏金猎人、'
        '坊市商贩、遗迹寻宝者、被追杀者、复仇者、隐世高人、采药人、魔兽驯养者、官府差役 等。\n\n'
        '【条目格式】严格如下（每条空行分隔）：\n'
        '【名称】原型名（如"落魄的云游散修"）\n'
        '【身份】职业/社会地位\n'
        '【性格】2-3词（如"谨慎、贪财、重诺"）\n'
        '【动机】他想要什么（驱动剧情，如"急需金币救治同伴""寻找失传功法"）\n'
        '【说话风格】一句话示例台词（用第二人称"你"，体现性格）\n'
        '【可能带来的冲突/机遇】剧本钩子（如"以低价兜售来路不明的丹药""求你护送过危险地带"）\n'
        '【实力参考】大致等阶（与玩家前期接触的居多，斗者~斗师为主）'
    )
    return sys, user


def prompt_tropes():
    """桥段模板：多分支剧本的叙事套路（不是完整故事，是骨架）。"""
    sys = (
        '你是游戏叙事设计师。编写「桥段模板」素材——即随机事件的叙事套路骨架。'
        '每个模板是一个独立条目，用【名称】开头。'
        '模板要支持「3个结局」的分支结构。只输出条目文本，不要前言后语、不要 markdown。'
    )
    user = (
        f'【世界观约束】\n{WORLD_SUMMARY}\n\n'
        '【任务】编写 12 个桥段模板，覆盖：奇遇赠予、仇家寻仇、神秘委托、宝物线索、'
        '道德抉择、拍卖纠纷、遗迹探险、魔兽袭击、救人/被救、赌斗、勒索、故人重逢 等。\n\n'
        '【条目格式】严格如下（每条空行分隔）：\n'
        '【名称】桥段名（如"坊市神秘丹贩"）\n'
        '【主题】奇遇/冲突/委托/抉择/探索（择一）\n'
        '【触发情境】何时何地自然发生（如"玩家进入坊市时"）\n'
        '【开场冲突】一句话：事件如何起头（NPC登场 + 抛出矛盾）\n'
        '【分支选项】列出2-3个玩家可选的选项，每个导向不同结局\n'
        '【结局A】向好结局（玩家得利），简述 + 可能的奖励类型(金币/物品/修为)\n'
        '【结局B】平淡结局（无得无失或小得失），简述\n'
        '【结局C】坏结局（玩家受损或遭遇战斗），简述 + 可能的惩罚(扣金币/战斗/损失)\n'
        '【叙事要点】写好这条桥段要注意的氛围/世界观质感（一句话）'
    )
    return sys, user


def prompt_regions():
    """地域风物：剧本舞台，让剧情贴合地点特色。"""
    sys = (
        '你是斗气大陆世界观设定师。编写「地域风物」素材——剧本事件发生的舞台背景。'
        '每个地域是独立条目，用【名称】开头。只输出条目文本，不要前言后语、不要 markdown。'
    )
    user = (
        f'【世界观约束】\n{WORLD_SUMMARY}\n\n'
        '【任务】编写 10 个地域风物条目，覆盖玩家可能途径的各类地点：'
        '沙漠商道、魔兽山脉外围、坊市、古城遗迹、密林深处、火山地带、'
        '冰雪荒原、湖畔小镇、黑角域边缘、险峻山隘 等。\n\n'
        '【条目格式】严格如下（每条空行分隔）：\n'
        '【名称】地名\n'
        '【地貌】一句话地形/气候\n'
        '【特色风物】此地独有的人文/自然景观（如"商队驼铃""地热温泉""废弃祭坛"）\n'
        '【常见魔兽】可能遭遇的魔兽类型（呼应魔兽图鉴，如"火属性一阶魔兽"）\n'
        '【常见机遇】在此地可能发生的奇遇类型（如"发现古修士遗物""采到珍稀药材"）\n'
        '【常见危险】在此地可能遭遇的风险（如"沙暴""劫匪""高阶魔兽"）\n'
        '【氛围词】3-5个关键词（用于剧本台词渲染，如"萧索、危机四伏、炽热"）'
    )
    return sys, user


SECTIONS = [
    ('势力档案', prompt_factions),
    ('NPC原型', prompt_npcs),
    ('桥段模板', prompt_tropes),
    ('地域风物', prompt_regions),
]


def main():
    dry = '--dry' in sys.argv
    print('=== 开始生成剧情素材（4 类，每类一次 LLM 调用）===')
    parts = []  # (标题, 文本)
    for title, prompt_fn in SECTIONS:
        print(f'\n>>> 生成 [{title}] ...')
        sys_p, user_p = prompt_fn()
        try:
            text = gen_section(sys_p, user_p)
            # 简单统计条目数
            n = text.count('【名称】')
            print(f'    完成：约 {n} 个条目，{len(text)} 字符')
            parts.append((title, text))
        except Exception as e:
            print(f'    ✗ 失败：{e}')
            parts.append((title, f'【生成失败：{e}】'))

    if dry:
        print('\n=== [DRY] 不写文件，预览如下 ===')
        for title, text in parts:
            print(f'\n{"="*20} {title} {"="*20}')
            print(text[:1500])
        return

    # 写入 docx
    doc = Document()
    doc.add_heading('剧情素材', level=0)
    doc.add_paragraph(
        '本文档为剧本生成专用的叙事素材库，由 LLM 生成。包含四类：势力档案、NPC原型、桥段模板、地域风物。'
        '供 rag_service 按条目切分检索，驱动多分支随机事件剧本的生成。'
    )
    for title, text in parts:
        doc.add_heading(title, level=1)
        for line in text.splitlines():
            line = line.rstrip()
            if line.strip():
                doc.add_paragraph(line)
        doc.add_paragraph('')  # 段落间空行，便于条目切分

    os.makedirs(RAG_DIR, exist_ok=True)
    doc.save(OUT_PATH)
    print(f'\n✅ 已写入：{OUT_PATH}')
    print('下一步：review 内容后，运行  python rag_service.py index  重建索引入库。')


if __name__ == '__main__':
    main()
