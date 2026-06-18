"""
生成材料图鉴docx + 更新item表description
"""
from docx import Document
from docx.shared import Pt
import json, pymysql, re

DB = {
    'host': 'os.environ.get("DB_HOST", "127.0.0.1")',
    'port': 3306,
    'user': 'root',
    'password': 'os.environ.get("DB_PASSWORD", "")',
    'database': 'dqdl',
    'charset': 'utf8mb4',
}

# 加载数据
with open('material_data.json', 'r', encoding='utf-8') as f:
    materials = json.load(f)

with open('material_descriptions.json', 'r', encoding='utf-8') as f:
    descriptions = json.load(f)

print(f"材料数据: {len(materials)} 条")
print(f"描述数据: {len(descriptions)} 条")

# 生成材料图鉴docx
print("\n=== 生成材料图鉴 ===")
mat_doc = Document()
style = mat_doc.styles['Normal']
style.font.size = Pt(11)
style.font.name = '宋体'

for m in sorted(materials, key=lambda x: x['item_id']):
    mat_doc.add_paragraph(f"【ID】{m['item_id']}")
    mat_doc.add_paragraph(f"【名称】{m['name']}")
    
    # 来源魔兽
    source_parts = []
    for s in m['source_mobs']:
        source_parts.append(f"{s['mob_id']} {s['name']}")
    mat_doc.add_paragraph(f"【来源】{', '.join(source_parts)}")
    
    mat_doc.add_paragraph(f"【类型】材料")
    
    # 稀有度：来源越多越常见
    source_count = len(m['source_mobs'])
    if source_count >= 5:
        rarity = '常见'
    elif source_count >= 2:
        rarity = '不常见'
    else:
        rarity = '稀有'
    mat_doc.add_paragraph(f"【稀有度】{rarity}")
    
    # AI生成的描述
    desc = descriptions.get(m['name'], f'由{m["source_mobs"][0]["name"] if m["source_mobs"] else "魔兽"}产出的炼制材料')
    mat_doc.add_paragraph(f"【描述】{desc}")
    
    mat_doc.add_paragraph('---')

mat_doc.save('../rag/材料图鉴.docx')
print(f"材料图鉴已保存: {len(materials)} 条")

# 更新item表description
print("\n=== 更新数据库 ===")
conn = pymysql.connect(**DB)
cur = conn.cursor()

updated = 0
for m in materials:
    desc = descriptions.get(m['name'], '')
    if desc:
        cur.execute(
            "UPDATE item SET description = %s WHERE item_id = %s",
            (desc, m['item_id'])
        )
        updated += 1

conn.commit()

# 验证
cur.execute("SELECT COUNT(*) FROM item WHERE type='材料' AND description IS NOT NULL AND description != ''")
print(f"材料有描述的: {cur.fetchone()[0]}")
cur.execute("SELECT item_id, name, description FROM item WHERE type='材料' LIMIT 5")
for row in cur.fetchall():
    print(f"  {row[0]} | {row[1]} | {row[2]}")

conn.close()
print("\n数据库更新完成!")
