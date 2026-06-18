"""分析魔兽图鉴docx的当前状态：掉落物和属性参考"""
from docx import Document
import re

doc = Document('../rag/魔兽图鉴.docx')

entries = []
current = {}
for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    if t.startswith('【ID】WB-'):
        if current and 'mob_id' in current:
            entries.append(current)
        current = {'mob_id': re.search(r'WB-\d+', t).group()}
    elif t.startswith('【名称】'):
        current['name'] = t.replace('【名称】', '').strip()
    elif t.startswith('【分类】'):
        current['element'] = t.replace('【分类】', '').strip()
    elif t.startswith('【品阶】'):
        current['tier'] = t.replace('【品阶】', '').strip()
    elif t.startswith('【掉落物】'):
        current['drops'] = t.replace('【掉落物】', '').strip()
    elif t.startswith('【属性参考】'):
        current['attrs'] = t.replace('【属性参考】', '').strip()
    elif t.startswith('【战力参考】'):
        current['power_ref'] = t.replace('【战力参考】', '').strip()
if current and 'mob_id' in current:
    entries.append(current)

print(f'总条目数: {len(entries)}')
print(f'有掉落物的: {sum(1 for e in entries if e.get("drops"))}')
print(f'无掉落物的: {sum(1 for e in entries if not e.get("drops"))}')
print(f'有属性参考的: {sum(1 for e in entries if e.get("attrs"))}')
print(f'无属性参考的: {sum(1 for e in entries if not e.get("attrs"))}')

# 分析掉落物格式
print('\n=== 掉落物格式样例(前10条) ===')
for e in entries[:10]:
    print(f'{e["mob_id"]} {e.get("name","")} | drops={e.get("drops","NONE")}')

# 分析属性参考格式
print('\n=== 属性参考格式样例(前10条) ===')
for e in entries[:10]:
    print(f'{e["mob_id"]} {e.get("name","")} | attrs={e.get("attrs","NONE")}')

# 收集所有掉落物中的item_id
import re
all_item_ids = set()
all_drop_names = []
for e in entries:
    d = e.get('drops', '')
    if d:
        ids = re.findall(r'【ID】([\w-]+)', d)
        all_item_ids.update(ids)
        # 提取掉落物名称(去掉【ID】xxx前缀后的内容)
        names = re.findall(r'【ID】[\w-]+\s*(.*?)(?:，|$)', d)
        all_drop_names.extend([n.strip() for n in names if n.strip()])

print(f'\n=== 掉落物统计 ===')
print(f'不同item_id数量: {len(all_item_ids)}')
print(f'mh开头(魔核): {sorted([i for i in all_item_ids if i.startswith("mh")])[:20]}')
print(f'cl开头(材料): {sorted([i for i in all_item_ids if i.startswith("cl")])[:20]}')
print(f'掉落物名称样例: {all_drop_names[:10]}')

# 检查材料图鉴
print('\n=== 材料图鉴 ===')
try:
    mat_doc = Document('../rag/材料图鉴.docx')
    mat_entries = []
    cur = {}
    for p in mat_doc.paragraphs:
        t = p.text.strip()
        if not t or t == '---':
            if cur and 'id' in cur:
                mat_entries.append(cur)
            cur = {}
            continue
        if t.startswith('【ID】'):
            cur['id'] = t.replace('【ID】', '').strip()
        elif t.startswith('【名称】'):
            cur['name'] = t.replace('【名称】', '').strip()
        elif t.startswith('【来源】'):
            cur['source'] = t.replace('【来源】', '').strip()
        elif t.startswith('【类型】'):
            cur['type'] = t.replace('【类型】', '').strip()
        elif t.startswith('【描述】') or t.startswith('【用途】'):
            cur['desc'] = t.replace('【描述】', '').replace('【用途】', '').strip()
    if cur and 'id' in cur:
        mat_entries.append(cur)

    print(f'材料图鉴条目数: {len(mat_entries)}')
    for m in mat_entries[:5]:
        print(f"  {m.get('id')} | {m.get('name','')} | src={m.get('source','')} | desc={m.get('desc','')}")
except Exception as ex:
    print(f'材料图鉴读取失败: {ex}')
