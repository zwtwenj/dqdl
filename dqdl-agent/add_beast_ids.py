"""为魔兽图鉴中每只魔兽插入唯一ID"""
from docx import Document
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy

DOC_PATH = r'rag\魔兽图鉴.docx'
OUT_PATH = r'rag\魔兽图鉴.docx'  # 直接覆盖原文件

doc = Document(DOC_PATH)

# 找每个魔兽条目的起始位置（【名称】行）
entries = []
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('【名称】'):
        entries.append(i)

print(f'找到 {len(entries)} 只魔兽')

body = doc.element.body
all_children = list(body)

# 从后往前插入ID，避免索引偏移
for idx, pos in enumerate(reversed(entries)):
    beast_num = len(entries) - idx  # 从 1 开始
    beast_id = f'WB-{beast_num:03d}'
    
    # 在【名称】段落前插入一个新段落作为ID
    name_elem = doc.paragraphs[pos]._element
    
    # 使用 python-docx API 创建段落
    new_p = parse_xml(f'<w:p {nsdecls("w")}><w:r><w:rPr></w:rPr><w:t xml:space="preserve">【ID】{beast_id}</w:t></w:r></w:p>')
    
    # 插入到name_elem之前
    name_elem.addprevious(new_p)

print('保存中...')
doc.save(OUT_PATH)
print(f'完成！已为 {len(entries)} 只魔兽添加ID，保存至 {OUT_PATH}')
