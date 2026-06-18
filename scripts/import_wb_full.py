"""解析魔兽图鉴全部303条，生成数据并直接插入数据库和材料图鉴"""
from docx import Document
from docx.shared import Pt
import json, re, pymysql, os

DB = {
    'host': 'os.environ.get("DB_HOST", "127.0.0.1")',
    'port': 3306,
    'user': 'root',
    'password': 'os.environ.get("DB_PASSWORD", "")',
    'database': 'dqdl',
    'charset': 'utf8mb4',
}

doc = Document('../rag/魔兽图鉴.docx')

# 逐段落解析每条魔兽
entries = []
current = {}
for p in doc.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    if t.startswith('【ID】WB-'):
        if current and 'mob_id' in current:
            entries.append(current)
        current = {'mob_id': re.search(r'WB-\d+', t).group()}
    elif t.startswith('【名称】'):
        current['name'] = t.replace('【名称】', '').strip()
    elif t.startswith('【分类】'):
        current['element'] = t.replace('【分类】', '').strip()
    elif t.startswith('【品阶】'):
        current['tier'] = t.replace('【品阶】', '').strip()
    elif t.startswith('【危险度】'):
        current['danger'] = t.replace('【危险度】', '').strip()
    elif t.startswith('【外观】'):
        current['appearance'] = t.replace('【外观】', '').strip()
    elif t.startswith('【体型】'):
        current['size'] = t.replace('【体型】', '').strip()
    elif t.startswith('【核心能力】'):
        current['abilities'] = t.replace('【核心能力】', '').strip()
    elif t.startswith('【弱点】'):
        current['weakness'] = t.replace('【弱点】', '').strip()
    elif t.startswith('【魔核】'):
        current['core'] = t.replace('【魔核】', '').strip()
    elif t.startswith('【掉落物】'):
        current['drops_raw'] = t.replace('【掉落物】', '').strip()
    elif t.startswith('【战力参考】'):
        current['power_ref'] = t.replace('【战力参考】', '').strip()
    elif t.startswith('【稀有度】'):
        current['rarity'] = t.replace('【稀有度】', '').strip()
    elif t.startswith('【性情】'):
        current['temper'] = t.replace('【性情】', '').strip()
if current and 'mob_id' in current:
    entries.append(current)

print(f'Parsed {len(entries)} entries')

# 品阶→属性模板
TIER_ATTRS = {
    '一阶': {'power': 8, 'intelligence': 5, 'quick': 10, 'stamina': 7},
    '二阶': {'power': 20, 'intelligence': 14, 'quick': 22, 'stamina': 18},
    '三阶': {'power': 40, 'intelligence': 30, 'quick': 38, 'stamina': 35},
    '四阶': {'power': 70, 'intelligence': 55, 'quick': 62, 'stamina': 65},
    '五阶': {'power': 120, 'intelligence': 100, 'quick': 105, 'stamina': 110},
}

POWER_MOD = {
    '一星斗师': 0.3, '二星斗师': 0.5, '三星斗师': 0.7, '四星斗师': 0.85,
    '五星斗师': 1.0, '六星斗师': 1.2, '七星斗师': 1.4,
}

MH_MAP = {
    ('火','一阶'): ('mh-h1-1', '劣质一阶火属性魔核'), ('火','二阶'): ('mh-h2-1', '普通二阶火属性魔核'), ('火','三阶'): ('mh-h3-1', '精良三阶火属性魔核'), ('火','四阶'): ('mh-h4-1', '稀有四阶火属性魔核'), ('火','五阶'): ('mh-h5-1', '传说五阶火属性魔核'),
    ('水','一阶'): ('mh-h1-2', '劣质一阶水属性魔核'), ('水','二阶'): ('mh-h2-2', '普通二阶水属性魔核'), ('水','三阶'): ('mh-h3-2', '精良三阶水属性魔核'), ('水','四阶'): ('mh-h4-2', '稀有四阶水属性魔核'), ('水','五阶'): ('mh-h5-2', '传说五阶水属性魔核'),
    ('风','一阶'): ('mh-h1-3', '劣质一阶风属性魔核'), ('风','二阶'): ('mh-h2-3', '普通二阶风属性魔核'), ('风','三阶'): ('mh-h3-3', '精良三阶风属性魔核'), ('风','四阶'): ('mh-h4-3', '稀有四阶风属性魔核'), ('风','五阶'): ('mh-h5-3', '传说五阶风属性魔核'),
    ('雷','一阶'): ('mh-h1-4', '劣质一阶雷属性魔核'), ('雷','二阶'): ('mh-h2-4', '普通二阶雷属性魔核'), ('雷','三阶'): ('mh-h3-4', '精良三阶雷属性魔核'), ('雷','四阶'): ('mh-h4-4', '稀有四阶雷属性魔核'), ('雷','五阶'): ('mh-h5-4', '传说五阶雷属性魔核'),
    ('冰','一阶'): ('mh-h1-5', '劣质一阶冰属性魔核'), ('冰','二阶'): ('mh-h2-5', '普通二阶冰属性魔核'), ('冰','三阶'): ('mh-h3-5', '精良三阶冰属性魔核'), ('冰','四阶'): ('mh-h4-5', '稀有四阶冰属性魔核'), ('冰','五阶'): ('mh-h5-5', '传说五阶冰属性魔核'),
    ('土','一阶'): ('mh-h1-6', '劣质一阶土属性魔核'), ('土','二阶'): ('mh-h2-6', '普通二阶土属性魔核'), ('土','三阶'): ('mh-h3-6', '精良三阶土属性魔核'), ('土','四阶'): ('mh-h4-6', '稀有四阶土属性魔核'), ('土','五阶'): ('mh-h5-6', '传说五阶土属性魔核'),
    ('木','一阶'): ('mh-h1-7', '劣质一阶木属性魔核'), ('木','二阶'): ('mh-h2-7', '普通二阶木属性魔核'), ('木','三阶'): ('mh-h3-7', '精良三阶木属性魔核'), ('木','四阶'): ('mh-h4-7', '稀有四阶木属性魔核'), ('木','五阶'): ('mh-h5-7', '传说五阶木属性魔核'),
    ('光','一阶'): ('mh-h1-8', '劣质一阶光属性魔核'), ('光','二阶'): ('mh-h2-8', '普通二阶光属性魔核'), ('光','三阶'): ('mh-h3-8', '精良三阶光属性魔核'), ('光','四阶'): ('mh-h4-8', '稀有四阶光属性魔核'), ('光','五阶'): ('mh-h5-8', '传说五阶光属性魔核'),
    ('暗','一阶'): ('mh-h1-9', '劣质一阶暗属性魔核'), ('暗','二阶'): ('mh-h2-9', '普通二阶暗属性魔核'), ('暗','三阶'): ('mh-h3-9', '精良三阶暗属性魔核'), ('暗','四阶'): ('mh-h4-9', '稀有四阶暗属性魔核'), ('暗','五阶'): ('mh-h5-9', '传说五阶暗属性魔核'),
    ('毒','一阶'): ('mh-h1-10', '劣质一阶毒属性魔核'), ('毒','二阶'): ('mh-h2-10', '普通二阶毒属性魔核'), ('毒','三阶'): ('mh-h3-10', '精良三阶毒属性魔核'), ('毒','四阶'): ('mh-h4-10', '稀有四阶毒属性魔核'), ('毒','五阶'): ('mh-h5-10', '传说五阶毒属性魔核'),
    ('斗','一阶'): ('mh-h1-11', '劣质一阶斗属性魔核'), ('斗','二阶'): ('mh-h2-11', '普通二阶斗属性魔核'), ('斗','三阶'): ('mh-h3-11', '精良三阶斗属性魔核'), ('斗','四阶'): ('mh-h4-11', '稀有四阶斗属性魔核'), ('斗','五阶'): ('mh-h5-11', '传说五阶斗属性魔核'),
}

def calc_attrs(entry):
    tier = entry.get('tier', '二阶')
    base = TIER_ATTRS.get(tier, TIER_ATTRS['二阶']).copy()
    power_ref = entry.get('power_ref', '三星斗师')
    mod = POWER_MOD.get(power_ref, 0.7)
    return {
        'power': round(base['power'] * mod),
        'intelligence': round(base['intelligence'] * mod),
        'quick': round(base['quick'] * mod),
        'stamina': round(base['stamina'] * mod),
    }

def build_drops(entry, material_id_map):
    """构建掉落JSON"""
    raw = entry.get('drops_raw', '')
    drops = []
    element = entry.get('element','')
    tier = entry.get('tier','')
    
    # 魔核
    mh_info = MH_MAP.get((element, tier))
    if mh_info:
        drops.append({'item_id': mh_info[0], 'name': mh_info[1], 'rate': 0.5, 'min': 1, 'max': 1, 'type': '公共'})
    
    # 材质掉落
    for item in raw.split('、'):
        item = item.strip()
        if not item:
            continue
        cl_id = material_id_map.get(item, 'unknown')
        drops.append({'item_id': cl_id, 'name': item, 'rate': 0.75, 'min': 1, 'max': 3, 'type': '专属'})
    return json.dumps(drops, ensure_ascii=False)

def build_description(entry):
    parts = []
    if entry.get('appearance'):
        parts.append(f"【外观】{entry['appearance']}")
    if entry.get('size'):
        parts.append(f"【体型】{entry['size']}")
    if entry.get('abilities'):
        parts.append(f"【核心能力】{entry['abilities']}")
    if entry.get('weakness'):
        parts.append(f"【弱点】{entry['weakness']}")
    return '\n'.join(parts)

# ============ 数据库连接 ============
conn = pymysql.connect(**DB)
cur = conn.cursor()

# 1. 清空mob表（已有的WB-001~WB-011将重新插入完整数据）
cur.execute("SET FOREIGN_KEY_CHECKS=0")
cur.execute("DELETE FROM mob")
cur.execute("SET FOREIGN_KEY_CHECKS=1")
print("Mob table cleared")

# 2. 收集所有掉落物名称并分配cl编号
# 先查已有cl item
cur.execute("SELECT item_id, name FROM item WHERE item_id LIKE 'cl-%'")
existing_cl = {row[1]: row[0] for row in cur.fetchall()}
print(f"Existing CL items: {len(existing_cl)}")

# 收集所有新材料
all_materials = set()
material_sources = {}  # name -> [mob_ids]
for e in entries:
    raw = e.get('drops_raw', '')
    if raw:
        for item in raw.split('、'):
            item = item.strip()
            if item:
                all_materials.add(item)
                if item not in material_sources:
                    material_sources[item] = []
                material_sources[item].append(e['mob_id'])

new_materials = {m for m in all_materials if m not in existing_cl}
print(f"New materials to create: {len(new_materials)}")

# 分配新cl编号
next_cl = max([int(v.replace('cl-','')) for v in existing_cl.values()] + [100]) + 1
material_id_map = dict(existing_cl)
for m in sorted(new_materials):
    material_id_map[m] = f'cl-{next_cl}'
    next_cl += 1

# 3. 插入新材料到item表
new_material_inserts = 0
for name in sorted(new_materials):
    cl_id = material_id_map[name]
    price = 200 + (abs(hash(name)) % 300)
    cur.execute(
        "INSERT INTO item (item_id, name, type, price) VALUES (%s, %s, 'material', %s)",
        (cl_id, name, price)
    )
    new_material_inserts += 1
print(f"Inserted {new_material_inserts} new material items")

# 4. 插入所有303条魔兽数据
mob_inserts = 0
for e in entries:
    mob_id = e['mob_id']
    name = e.get('name', mob_id)
    attrs = calc_attrs(e)
    attribute = e.get('element', '')[:8]
    description = build_description(e)
    drops = build_drops(e, material_id_map)
    
    cur.execute(
        "INSERT INTO mob (mob_id, name, power, intelligence, quick, stamina, attribute, description, drops) "
        "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
        (mob_id, name, attrs['power'], attrs['intelligence'], attrs['quick'], attrs['stamina'], attribute, description, drops)
    )
    mob_inserts += 1
print(f"Inserted {mob_inserts} mob entries")

conn.commit()
conn.close()
print("Database committed!")

# 5. 生成材料图鉴docx（含来源魔兽ID）
mat_doc = Document()
style = mat_doc.styles['Normal']
style.font.size = Pt(11)
style.font.name = '宋体'

for name, cl_id in sorted(material_id_map.items(), key=lambda x: x[1]):
    mat_doc.add_paragraph(f'【ID】{cl_id}')
    mat_doc.add_paragraph(f'【名称】{name}')
    sources = material_sources.get(name, [])
    source_str = '、'.join(sources) if sources else '未知'
    mat_doc.add_paragraph(f'【来源】{source_str}')
    mat_doc.add_paragraph(f'【类型】材料')
    # 估计稀有度：来源越多越常见
    if len(sources) >= 5:
        rarity = '常见'
    elif len(sources) >= 2:
        rarity = '不常见'
    else:
        rarity = '稀有'
    mat_doc.add_paragraph(f'【稀有度】{rarity}')
    mat_doc.add_paragraph(f'【用途】可用于炼制丹药或打造装备。')
    mat_doc.add_paragraph('---')

# 注意：如果材料图鉴已存在，用新版本覆盖
mat_doc.save('../rag/材料图鉴.docx')
print(f'Generated 材料图鉴.docx with {len(material_id_map)} materials')

print('All done!')
