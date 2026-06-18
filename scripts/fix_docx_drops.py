"""修复魔兽图鉴.docx的【掉落物】格式，统一为纯材料名+、分隔"""
from docx import Document
from docx.shared import Pt
import re, copy

doc = Document('../rag/魔兽图鉴.docx')

# 记录需要修复的段落
fixes = []

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t.startswith('【掉落物】'):
        continue

    raw = t.replace('【掉落物】', '').strip()
    if not raw:
        continue

    # 策略：从原文中提取所有独立的材料名称
    # 支持三种格式：
    # 1. "【ID】mh-xxx 描述。【ID】cl-xxx 焰鳞片（描述）" → 提取 cl-xxx 后的名称
    # 2. "冰晶翼×2，冰丝触须×1" → 拆分为 冰晶翼、冰丝触须
    # 3. "霜纹蛙皮（描述）、冰舌黏液（描述）。" → 拆分为 霜纹蛙皮、冰舌黏液

    items = []

    # 先去掉【ID】开头的魔核描述段落（整段可能是魔核，后面跟着cl材料）
    # 拆分：先按 ， 和 、 拆成段
    segments = re.split(r'[，、]', raw)

    for seg in segments:
        seg = seg.strip().rstrip('。')
        if not seg:
            continue

        # 跳过纯魔核描述（含【ID】mh- 的）
        if '【ID】mh-' in seg or '【ID】mh' in seg:
            # 但如果后面有 cl 材料，也尝试提取
            # 格式: "【ID】mh-h1-1 ...。【ID】cl-100 焰鳞片（...）"
            cl_matches = re.findall(r'【ID】cl-\d+\s*([^(（\u3001\uff0c]+)', seg)
            for m in cl_matches:
                name = m.strip()
                if name and name not in items:
                    items.append(name)
            continue

        # 跳过纯【ID】cl-xxx 描述行（已被上面处理）
        if '【ID】cl-' in seg:
            # 提取材料名
            m = re.search(r'【ID】cl-\d+\s*([^(（]+)', seg)
            if m:
                name = m.group(1).strip()
                if name and name not in items:
                    items.append(name)
            continue

        # 普通材料名：去掉 ×N 数量标记，去掉 （描述）
        name = re.sub(r'\u00d7\d+', '', seg)  # 去 ×N
        name = re.sub(r'[（(][^）)]*[）)]', '', name)  # 去括号描述
        name = name.strip()

        if name and name not in items:
            items.append(name)

    if items:
        new_text = '【掉落物】' + '、'.join(items)
        # 只在内容变了才改
        if new_text != t:
            # 修改段落
            # 需要清空原来的runs，写入新文本
            for run in p.runs:
                run.text = ''
            if p.runs:
                p.runs[0].text = new_text
            else:
                p.add_run(new_text)
            fixes.append((i, t[:60], new_text[:60]))

print(f'修复了 {len(fixes)} 条【掉落物】')
for idx, old, new in fixes[:10]:
    print(f'  [{idx}] {old} → {new}')
if len(fixes) > 10:
    print(f'  ... 等共 {len(fixes)} 条')

doc.save('../rag/魔兽图鉴.docx')
print('\n已保存 ../rag/魔兽图鉴.docx')
