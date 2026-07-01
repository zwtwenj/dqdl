# -*- coding: utf-8 -*-
"""
炼丹图鉴生成器：读取 alchemy_data.json，生成 rag/草药图鉴.docx 与 rag/丹药图鉴.docx。
用法: python scripts/build_alchemy_docx.py
"""
import os, json, sys

try:
    from docx import Document
except ImportError:
    print('请先 pip install python-docx', file=sys.stderr)
    sys.exit(1)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_PATH = os.path.join(ROOT, 'scripts', 'alchemy_data.json')
RAG_DIR = os.path.join(ROOT, 'rag')
HERB_DOCX = os.path.join(RAG_DIR, '草药图鉴.docx')
PILL_DOCX = os.path.join(RAG_DIR, '丹药图鉴.docx')

ELEMENTS = ['金', '木', '水', '火', '土', '雷', '风']


def fmt_element(el: dict) -> str:
    """{木:10,火:5} -> '木10 火5'，按固定元素顺序输出。"""
    return ' '.join(f'{e}{el[e]}' for e in ELEMENTS if e in el)


def fmt_recipe(required: dict, tolerance: dict) -> str:
    """{木:20,火:10},tol{木:4,火:2} -> '木20±4 火10±2'"""
    return ' '.join(
        f'{e}{required[e]}±{tolerance.get(e, 0)}'
        for e in ELEMENTS if e in required
    )


def write_herb_docx(data):
    doc = Document()
    doc.add_heading('草药图鉴', level=0)
    doc.add_paragraph('斗气大陆常见草药，每株含元素能量（金木水火土雷风）。元素能量可用于炼丹配方匹配。')
    for h in data['herbs']:
        lines = [
            f"【名称】{h['name']}",
            f"【品阶】{h['tier']}阶",
            f"【类别】{h['category']}",
            f"【元素】{fmt_element(h['element'])}",
            f"【稀有度】{h['rarity']}",
            f"【栖息地】{h['habitat']}",
            f"【外观】{h['appearance']}",
            f"【核心功效】{h['effect']}",
            f"【市场价】{h['price']}金币/株",
        ]
        for ln in lines:
            doc.add_paragraph(ln)
        doc.add_paragraph('---')
    doc.save(HERB_DOCX)
    print(f'✅ 写入 {HERB_DOCX}  共 {len(data["herbs"])} 条草药')


def write_pill_docx(data):
    # 丹方格式说明给 RAG
    furnace_tier_name = {1: '黄阶', 2: '玄阶', 3: '地阶', 4: '天阶'}
    recipe_by_out = {r['output_item_id']: r for r in data['recipes']}
    doc = Document()
    doc.add_heading('丹药图鉴', level=0)
    doc.add_paragraph(
        '斗气大陆丹药图鉴。每味丹药含【丹方】= 所需元素能量目标与公差（如 木20±4 火10±2 表示'
        '投入材料的木能量落在16~24、火能量落在8~12即算匹配成功）。丹炉品阶决定可炼丹药品阶与'
        '单元素能量上限：黄阶100/玄阶300/地阶800/天阶2000。'
    )
    for p in data['pills']:
        r = recipe_by_out.get(p['item_id'])
        lines = [
            f"【名称】{p['name']}",
            f"【品阶】{p['tier']}阶",
            f"【类别】{p['category']}",
        ]
        if r:
            minf = furnace_tier_name.get(r['min_furnace_tier'], '黄阶')
            lines.append(f"【丹方】{fmt_recipe(r['required'], r['tolerance'])}")
            lines.append(f"【最低丹炉】{minf}丹炉")
        lines.append(f"【核心效果】{p['desc']}")
        lines.append(f"【市价】{p['price']}金币")
        lines.append(f"【适用境界】{'斗之气' if p['tier'] == 1 else '斗者' if p['tier'] == 2 else '斗师及以上'}")
        for ln in lines:
            doc.add_paragraph(ln)
        doc.add_paragraph('---')
    # 丹炉附录
    doc.add_heading('丹炉附录', level=1)
    for f in data['furnaces']:
        doc.add_paragraph(
            f"【名称】{f['name']}\n"
            f"【品阶】{furnace_tier_name.get(f['tier'], '?')}阶\n"
            f"【投放槽位】{f['spec']['slots']}种\n"
            f"【单元素能量上限】{f['spec']['cap']}\n"
            f"【耐久】{f['spec']['max_durability']}\n"
            f"【说明】{f['desc']}\n---"
        )
    doc.save(PILL_DOCX)
    print(f'✅ 写入 {PILL_DOCX}  共 {len(data["pills"])} 条丹药 + {len(data["furnaces"])} 个丹炉')


def main():
    with open(DATA_PATH, encoding='utf-8') as fp:
        data = json.load(fp)
    os.makedirs(RAG_DIR, exist_ok=True)
    write_herb_docx(data)
    write_pill_docx(data)


if __name__ == '__main__':
    main()
